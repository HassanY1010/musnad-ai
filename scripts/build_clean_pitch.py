import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_clean_pitch():
    template_path = 'E:/basira/musnad-ai/1.pptx'
    out_path = 'E:/basira/musnad-ai/MUSNAD-Pitch-Deck-Final.pptx'
    
    prs = Presentation(template_path)
    
    # 1. Remove all 31 template placeholder slides
    num_to_remove = len(prs.slides)
    for _ in range(num_to_remove):
        rId = prs.slides._sldIdLst[0]
        prs.part.drop_rel(rId.rId)
        prs.slides._sldIdLst.remove(rId)
        
    print(f"Removed {num_to_remove} placeholder slides. Slides remaining: {len(prs.slides)}")
    
    # Helper to add standard content slide using layout 1
    def add_slide(title, lines):
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        
        # Title
        if len(slide.placeholders) >= 1:
            title_shape = slide.placeholders[0]
            title_shape.text = title
            for p in title_shape.text_frame.paragraphs:
                p.alignment = PP_ALIGN.RIGHT
                p.font.language_id = 1025
                p.font.bold = True
                
        # Body
        if len(slide.placeholders) >= 2:
            body_shape = slide.placeholders[1]
            tf = body_shape.text_frame
            tf.clear()
            for line in lines:
                p = tf.add_paragraph()
                p.text = line
                p.font.language_id = 1025
                if line and line[0].encode('utf-8', 'ignore').isalpha():
                    p.alignment = PP_ALIGN.LEFT
                else:
                    p.alignment = PP_ALIGN.RIGHT
        return slide

    # Slide 1: Cover Title
    add_slide("مُسنَد | MUSNAD AI", [
        "محرك ذكي للتحقق من المحتوى الإسلامي وتتبع الأدلة الشرعية",
        " ",
        "المسار: أدوات المعرفة والتحقق لتمكين المعرفين بالإسلام",
        "تحدي الذكاء الاصطناعي في خدمة المحتوى الإسلامي 2026",
        " ",
        "المطور: حسن محمد بارفعه"
    ])

    # Slide 2: Problem
    add_slide("تحديات الموثوقية والمصداقية (المشكلة)", [
        "• فوضى رقمية: انتشار متسارع للاقتباسات والأحاديث الضعيفة والمكذوبة دون سند.",
        "• هلوسة تقنية: نماذج الذكاء الاصطناعي التوليدي تختلق نصوصاً وأسانيد لا وجود لها.",
        "• فتاوى آلية غير منضبطة: توليد أحكام فقهية قطعية في مسائل خلافية واجتهادية.",
        "• استنزاف الوقت: معاناة الباحثين والمترجمين في التثبت والتخريج الدقيق من المراجع المعتمدة."
    ])

    # Slide 3: Solution
    add_slide("الحل المبتكر: مُسنَد (Evidence Engine)", [
        "• محرك تدقيق حتمي مغلق (Deterministic Verification Engine).",
        "• تفكيك النصوص آلياً إلى ادعاءات ذرية قابلة للقياس والإسناد المصدري.",
        "• التحقق الحصري من قاعدة معرفة موثقة خالية من الهلوسة بنسبة 100%.",
        "• الامتناع الآمن (Safe Abstention) والإحالة للمختصين بدلاً من الفتوى التلقائية.",
        "• واجهة استعلام متقدمة تعرض تباين الروايات والفروق اللفظية بشفافية كاملة."
    ])

    # Slide 4: Verification Pipeline
    add_slide("دورة التحقق والتدقيق (Workflow)", [
        "1. استقبال المحتوى: معالجة النصوص والاقتباسات والادعاءات الدينية.",
        "2. استخراج الادعاءات: تجزئة المحتوى وتصنيفه (آية، حديث، قول، مسألة فقهية).",
        "3. التطبيع اللغوي العربي: معالجة الهمزات، حذف التشكيل والتطويل، لضمان أعلى دقة.",
        "4. الاسترجاع الهجين (Hybrid RAG): مطابقة نصية دقيقة وبحث معجمي BM25 فائق السرعة.",
        "5. التحقق بقواعد حتمية: مقارنة الادعاء بالشاهد المعتمد مع منع الـ LLM من القرار.",
        "6. المخرجات وسجل التدقيق: عرض الحكم، التخريج، الشاهد، وبصمة التشفير (SHA-256)."
    ])

    # Slide 5: Innovation & Reliability
    add_slide("الميزة التنافسية والسلامة العلمية", [
        "• قاعدة (Claim → Evidence → Deterministic Status): القرار برمجي قطعي 100%.",
        "• سياسة انعدام الفتوى (Zero-Fatwa Policy): الإحالة التلقائية في المسائل المعاصرة.",
        "• منع الهلوسة بنسبة 100%: امتناع كامل عند غياب الشاهد المعتمد (Safe Abstention).",
        "• تحليل الفروق اللفظية (Textual Difference): رصد اختلاف الألفاظ والروايات بدقة.",
        "• سجل تدقيق مشفر بسلسلة كتل (Cryptographic Audit Log Chain) لحفظ النزاهة.",
        "• مناعة متقدمة ضد هجمات التلاعب وحقن الأوامر (Prompt Injection Defense)."
    ])

    # Slide 6: Architecture
    add_slide("البنية التقنية والمعمارية (Architecture)", [
        "• الواجهة الأمامية: React 18 + TypeScript + Vite + Tailwind CSS (تصميم RTL فائق الجودة).",
        "• الواجهة البرمجية (API): FastAPI + Python 3.11 مع توثيق OpenAPI كامل.",
        "• محرك البحث والاسترجاع: Hybrid BM25Okapi + Substring Matching + RRF (k=60).",
        "• قاعدة البيانات: SQLite خفيفة عالية السرعة (مع توافق كامل مع pgvector في PostgreSQL).",
        "• المحتوى المفهرس: 48,475 سجلاً محققاً يشمل القرآن، الصحاح، السنن، والتفاسير.",
        "• الأداء: زمن استجابة متناهي الصغر (متوسط زمن التحقق ~4.5 مللي ثانية)."
    ])

    # Slide 7: Live Scenario
    add_slide("نموذج حي للتحقق (Live Verification Demo)", [
        "• النص المدخل: «إنما الأعمال بالنيات وإنما لكل امرئ ما نوى...»",
        "• النتيجة الآلية: مدعوم بدليل قطعي بمطابقة تامة (Supported - Exact Match).",
        "• الشاهد المعتمد: حديث أمير المؤمنين عمر بن الخطاب رضي الله عنه.",
        "• التوثيق الدقيق: صحيح البخاري (كتاب بدء الوحي - حديث 1) / الأربعون النووية (حديث 1).",
        "• حالة ادعاء مكذوب: امتناع كامل وتصنيف (Insufficient Evidence) دون أي اختلاق.",
        "• حالة مسألة فقهية: إحالة للمختص (Specialist Referral) لحماية الضوابط الشرعية."
    ])

    # Slide 8: Target Audience & Impact
    add_slide("الفئات المستهدفة والأثر المستدام", [
        "• المعرفون بالإسلام والجهات الدعوية: تقديم محتوى مترجم وموثق ببرهان قطعي.",
        "• الباحثون والمترجمون: اختصار ساعات طويلة من البحث وتخريج النصوص الشرعية.",
        "• منصات المحتوى الإسلامي: دمج محرك التدقيق عبر API كصمام أمان رقمي.",
        " ",
        "• الأثر المستدام:",
        "تعزيز الثقة في الفضاء الرقمي الإسلامي، محاصرة الشائعات والنسب غير الموثوقة، وتمكين التحول التقني الرشيد وفق مستهدفات رؤية المملكة 2030."
    ])

    # Slide 9: Conclusion
    add_slide("الخاتمة", [
        "مُسنَد — MUSNAD AI",
        "«محرك ذكي.. يحول الشبهة إلى تدقيق، والادعاء إلى برهان»",
        " ",
        "المشروع متكامل وجاهز للتشغيل والتقييم الحي (100% Offline Demo Ready)",
        " ",
        "المطور: حسن محمد بارفعه",
        "تحدي الذكاء الاصطناعي في خدمة المحتوى الإسلامي"
    ])

    prs.save(out_path)
    # Also overwrite MUSNAD-Template-Applied.pptx so any standard link works
    prs.save('E:/basira/musnad-ai/MUSNAD-Template-Applied.pptx')
    print(f"Clean pitch deck saved to {out_path} and MUSNAD-Template-Applied.pptx with exactly {len(prs.slides)} slides.")

if __name__ == '__main__':
    create_clean_pitch()
