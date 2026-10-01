import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def add_title_slide(prs, title_text, subtitle_text):
    # استخدام التخطيط الأول (غالباً شريحة العنوان)
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    
    # محاولة العثور على مربعات النص وتعبئتها
    if len(slide.placeholders) >= 1:
        title_shape = slide.placeholders[0]
        title_shape.text = title_text
        for p in title_shape.text_frame.paragraphs:
            p.alignment = PP_ALIGN.RIGHT
            p.font.language_id = 1025
            
    if len(slide.placeholders) >= 2:
        subtitle_shape = slide.placeholders[1]
        subtitle_shape.text = subtitle_text
        for p in subtitle_shape.text_frame.paragraphs:
            p.alignment = PP_ALIGN.RIGHT
            p.font.language_id = 1025

def add_content_slide(prs, title, content_lines):
    # استخدام التخطيط الثاني (غالباً عنوان ومحتوى)
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    
    if len(slide.placeholders) >= 1:
        title_shape = slide.placeholders[0]
        title_shape.text = title
        for p in title_shape.text_frame.paragraphs:
            p.alignment = PP_ALIGN.RIGHT
            p.font.language_id = 1025
            
    if len(slide.placeholders) >= 2:
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.clear()
        
        for line in content_lines:
            p = tf.add_paragraph()
            p.text = line
            p.font.language_id = 1025
            
            # محاذاة لليسار للتقنيات الإنجليزية
            if line and line[0].encode('utf-8', 'ignore').isalpha():
                p.alignment = PP_ALIGN.LEFT
            else:
                p.alignment = PP_ALIGN.RIGHT

print("Opening template 1.pptx...")
prs = Presentation('E:/basira/musnad-ai/1.pptx')

# --- الشريحة 1: الغلاف ---
add_title_slide(prs, "مُسنَد | MUSNAD AI", 
    "محرك ذكي للتحقق من المحتوى الإسلامي وتتبع الأدلة الشرعية\n\n"
    "مسار التحدي: أدوات المعرفة والتحقق لتمكين المعرفين بالإسلام\n"
    "المطور: حسن محمد بارفعه"
)

# --- الشريحة 2: المشكلة ---
add_content_slide(prs, "تحديات الموثوقية والمصداقية", [
    "• فوضى رقمية: انتشار سريع للاقتباسات والأحاديث غير الموثقة.",
    "• هلوسة تقنية: اختلاق نماذج الذكاء الاصطناعي العامة لنصوص شرعية لا أصل لها.",
    "• استنزاف الوقت: الباحثون يواجهون صعوبة بالغة في التحقق الفوري وتخريج الأحاديث للوصول للمصدر."
])

# --- الشريحة 3: الحل ---
add_content_slide(prs, "مُسنَد: محرك تدقيق حتمي", [
    "• تحويل النصوص العشوائية أو المحتوى إلى (ادعاءات) قابلة للقياس.",
    "• بحث دقيق في قواعد بيانات شرعية موثوقة ومغلقة حصرياً.",
    "• تحقق صارم بقواعد حتمية تمنع التأليف والهلوسة تماماً.",
    "• عرض النتيجة النهائية بشفافية مع إرفاق المصدر وتخريج النص."
])

# --- الشريحة 4: آلية العمل ---
add_content_slide(prs, "دورة التحقق الفوري", [
    "1. إدخال المحتوى: استقبال النص أو الاقتباس.",
    "2. استخراج الادعاءات: فهم النص وتفكيكه إلى جمل للقياس.",
    "3. تطبيع العربية: توحيد الهمزات وحذف التشكيل لضمان دقة الاسترجاع.",
    "4. البحث الهجين (Retrieval): استخراج الأدلة من بين عشرات الآلاف من النصوص.",
    "5. التحقق الحتمي: مقارنة الادعاء بالدليل حرفياً ودلالياً.",
    "6. النتيجة والمصدر: إصدار الحكم النهائي بوضوح مع توثيق المرجع."
])

# --- الشريحة 5: الابتكار والموثوقية ---
add_content_slide(prs, "الابتكار والميزة التنافسية", [
    "• الاسترجاع الهجين (Hybrid Retrieval): دمج البحث النصي والدلالي لضمان الشمولية.",
    "• التحقق الحتمي (Deterministic Verification): نظام مقيد بقواعد تمنعه من الفتوى.",
    "• تتبع الأدلة (Evidence Traceability): ربط كل قرار بدليله المرجعي.",
    "• انعدام الهلوسة (Zero-Hallucination): النظام يعتمد على الأدلة فقط.",
    "• مبدأ انعدام الدليل (Abstention): إذا لم نجد دليلاً، لا يتم إصدار نتيجة مؤكدة.",
    "• الحماية من التلاعب (Prompt Injection Security)."
])

# --- الشريحة 6: البنية التقنية ---
add_content_slide(prs, "البنية التقنية والمعمارية", [
    "Frontend: React + TypeScript",
    "Backend: FastAPI + Python",
    "NLP: Claim Extraction / Arabic Processing Engine",
    "Retrieval: Hybrid Retrieval (BM25 + Dense Retrieval + RRF)",
    "Database: SQLite Vector + In-Memory O(1) Speed",
    "Core Engine: Deterministic Verification Engine (RAG Architecture)"
])

# --- الشريحة 7: مثال عملي ---
add_content_slide(prs, "مثال عملي للتحقق", [
    "• الادعاء المدخل: «إنما الأعمال بالنيات»",
    "• الأدلة المسترجعة: «عمر بن الخطاب رضي الله عنه قَالَ: سَمِعْتُ رَسُولَ اللَّهِ...»",
    "• المصدر الموثق: صحيح البخاري (حديث 1) / الأربعون النووية",
    "• حالة الادعاء: مدعوم بالدليل ✅",
    "• سبب النتيجة: تطابق تام في اللفظ والمعنى مع المصادر.",
    "(يرجى إدراج Screenshot لواجهة النظام الحقيقية هنا)"
])

# --- الشريحة 8: الأثر المستدام ---
add_content_slide(prs, "الأثر المستدام والفئات المستفيدة", [
    "• الجهات المستفيدة:",
    "  - الدعاة والباحثون الشرعيون الساعون للدقة المتناهية.",
    "  - صناع المحتوى الإسلامي والمؤسسات والمنصات الدعوية.",
    "  - المطورون الراغبون في دمج ميزة التدقيق عبر (API).",
    " ",
    "• الأثر المستدام:",
    "تسريع عملية التوثيق والتحقق بشكل جذري، رفع مستوى الشفافية، وتقليص المحتوى المضلل في الفضاء الرقمي."
])

# --- الشريحة 9: الخاتمة ---
add_content_slide(prs, "الخاتمة", [
    "مُسنَد — MUSNAD AI",
    " ",
    "«من ادعاء إلى دليل.. ومن دليل إلى مصدر قابل للتحقق»",
    " ",
    "المطور: حسن محمد بارفعه"
])

print("Saving final file...")
prs.save('E:/basira/musnad-ai/MUSNAD-Template-Applied.pptx')
print("Successfully saved to MUSNAD-Template-Applied.pptx")
